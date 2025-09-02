# -*- coding: utf-8 -*-

import openpyxl
from datetime import datetime, timedelta, date
import psycopg2
import ttkbootstrap as tb
from ttkbootstrap.constants import *
from ttkbootstrap import DateEntry
from tkinter import filedialog, messagebox, Toplevel, END, BOTH, YES, X, Y, LEFT, RIGHT, VERTICAL, HORIZONTAL, BOTTOM, CENTER, W, E

from config import LANGUAGES
from database import get_db_connection, release_db_connection

class ViewAppointmentsWindow(Toplevel):
    def __init__(self, master, db_config):
        super().__init__(master)
        self.master = master
        self.db_config = db_config
        self.title("Relatório de Produção (PCP)")
        self.geometry("1400x850")
        self.grab_set()

        self.kpi_labels = {}
        self.current_page = 1
        self.results_per_page = 100

        self.create_widgets()
        self.load_filter_options()
        self.apply_all_filters()

    def get_string(self, key, **kwargs):
        return self.master.get_string(key, **kwargs)

    def get_db_connection(self):
        return get_db_connection()

    def create_widgets(self):
        main_frame = tb.Frame(self, padding=10)
        main_frame.pack(fill=BOTH, expand=True)

        kpi_frame = tb.LabelFrame(main_frame, text="Indicadores Chave (do Histórico de Produção)", bootstyle=SUCCESS, padding=10)
        kpi_frame.pack(fill=X, pady=(0, 10))

        kpi_defs = {
            "qtd_produzida": "Qtd. Total Produzida",
            "giros_rodados": "Giros Rodados",
            "tempo_producao": "Tempo Total de Produção",
            "tempo_parada": "Tempo Total de Parada",
            "total_perdas": "Total de Perdas (Prod.)"
        }
        for i, (key, text) in enumerate(kpi_defs.items()):
            frame = tb.Frame(kpi_frame)
            frame.pack(side=LEFT, padx=20, pady=5, fill=X, expand=True)
            tb.Label(frame, text=text, font=("-weight bold")).pack()
            kpi_label = tb.Label(frame, text="-", font=("", 12))
            kpi_label.pack()
            self.kpi_labels[key] = kpi_label

        filter_frame = tb.LabelFrame(main_frame, text="Filtros do Relatório", bootstyle=PRIMARY, padding=10)
        filter_frame.pack(fill=X, pady=5)
        
        tb.Label(filter_frame, text="Data Início:").grid(row=0, column=0, padx=(0, 5), pady=5, sticky=W)
        self.date_start_entry = DateEntry(filter_frame, bootstyle=INFO, dateformat='%d/%m/%Y')
        self.date_start_entry.grid(row=0, column=1, padx=5, pady=5)
        tb.Label(filter_frame, text="Data Fim:").grid(row=0, column=2, padx=(10, 5), pady=5, sticky=W)
        self.date_end_entry = DateEntry(filter_frame, bootstyle=INFO, dateformat='%d/%m/%Y')
        self.date_end_entry.grid(row=0, column=3, padx=5, pady=5)
        tb.Label(filter_frame, text="Status Ordem:").grid(row=0, column=4, padx=(15, 5), pady=5, sticky=W)
        self.status_filter = tb.Combobox(filter_frame, state="readonly", values=["Todos", "Em Aberto", "Concluído"])
        self.status_filter.grid(row=0, column=5, padx=5, pady=5)
        self.status_filter.set("Todos")
        tb.Label(filter_frame, text="Cliente:").grid(row=1, column=0, padx=(0, 5), pady=5, sticky=W)
        self.client_filter = tb.Combobox(filter_frame, state="readonly")
        self.client_filter.grid(row=1, column=1, padx=5, pady=5)
        tb.Label(filter_frame, text="WO:").grid(row=1, column=2, padx=(10, 5), pady=5, sticky=W)
        self.wo_filter = tb.Combobox(filter_frame, state="readonly")
        self.wo_filter.grid(row=1, column=3, padx=5, pady=5)
        buttons_filter_frame = tb.Frame(filter_frame)
        buttons_filter_frame.grid(row=1, column=4, columnspan=2, padx=(15, 5), pady=5, sticky=E)
        tb.Button(buttons_filter_frame, text="Aplicar Filtros", command=self.apply_all_filters, bootstyle=SUCCESS).pack(side=LEFT, padx=10)
        tb.Button(buttons_filter_frame, text="Limpar Filtros", command=self.clear_filters, bootstyle=WARNING).pack(side=LEFT)

        tab_historico = tb.Frame(main_frame, padding=10)
        tab_historico.pack(fill=BOTH, expand=True, pady = 10)

        action_frame_hist = tb.LabelFrame(tab_historico, text="Ações do Histórico", bootstyle=INFO, padding=10)
        action_frame_hist.pack(fill=X, pady=(0,10))

        tb.Button(action_frame_hist, text="Abrir Gestão PCP (Ordens Pendentes)", command=self.open_pcp_window, bootstyle="primary-outline").pack(side=LEFT, padx=5)


        tb.Button(action_frame_hist, text="Ver Detalhes das Paradas", command=self.view_stops, bootstyle="secondary-outline").pack(side=LEFT, padx=5)
        tb.Button(action_frame_hist, text="Exportar Histórico para Excel", command=self.export_to_xlsx, bootstyle="success-outline").pack(side=RIGHT, padx=5)

        tree_frame_hist = tb.Frame(tab_historico)
        tree_frame_hist.pack(fill=BOTH, expand=True)

        self.cols_hist = (
            "apontamento_id", "wo", "cliente", "servico_desc", "servico_status", "data_apont", 
            "horainicio", "horafim", "impressor", "equipamento", "giros_rodados", "qtd_produzida", "perdas_producao"
        )
        self.headers_hist = [
            "ID Apont.", "WO", "Cliente", "Serviço", "Status Serviço", "Data", 
            "Início Prod.", "Fim Prod.", "Impressor", "Equipamento", "Giros Rodados", "Qtd. Produzida", "Perdas"
        ]
        
        self.tree_hist = tb.Treeview(tree_frame_hist, columns=self.cols_hist, show="headings", bootstyle=PRIMARY)
        for col, header in zip(self.cols_hist, self.headers_hist):
            self.tree_hist.heading(col, text=header)
            self.tree_hist.column(col, anchor=W, width=120)
        self.tree_hist.column("apontamento_id", width=0, stretch=False)
        self.tree_hist.pack(side=LEFT, fill=BOTH, expand=True)
        tb.Scrollbar(tree_frame_hist, orient=VERTICAL, command=self.tree_hist.yview).pack(side=RIGHT, fill=Y)

        pagination_frame = tb.Frame(tab_historico)
        pagination_frame.pack(fill=X, pady=(5,0))
        self.prev_page_button = tb.Button(pagination_frame, text="< Página Anterior", command=self.prev_page, bootstyle="secondary")
        self.prev_page_button.pack(side=LEFT)
        self.page_label = tb.Label(pagination_frame, text=f"Página {self.current_page}")
        self.page_label.pack(side=LEFT, padx=10)
        self.next_page_button = tb.Button(pagination_frame, text="Próxima Página >", command=self.next_page, bootstyle="secondary")
        self.next_page_button.pack(side=LEFT)

    def apply_all_filters(self):
        self.load_production_history()

    def load_filter_options(self):
        conn = self.get_db_connection()
        if not conn: return
        try:
            with conn.cursor() as cur:
                cur.execute("SELECT DISTINCT cliente FROM ordem_producao WHERE cliente IS NOT NULL ORDER BY cliente")
                clients = ["Todos"] + [row[0] for row in cur.fetchall()]
                self.client_filter['values'] = clients
                self.client_filter.set("Todos")

                cur.execute("SELECT DISTINCT numero_wo FROM ordem_producao ORDER BY numero_wo")
                wos = ["Todas"] + [row[0] for row in cur.fetchall()]
                self.wo_filter['values'] = wos
                self.wo_filter.set("Todas")
        except (psycopg2.Error, ValueError) as e:
            messagebox.showerror("Erro ao carregar filtros", f"Não foi possível carregar as opções de filtro: {e}", parent=self)
        finally:
            if conn: conn.close()
            
    def load_production_history(self):
        for item in self.tree_hist.get_children():
            self.tree_hist.delete(item)
        self.page_label.config(text=f"Página {self.current_page}")

        conn = self.get_db_connection()
        if not conn: return
        
        
        try:
            with conn.cursor() as cur:
                query = """
                    SELECT
                        ap.id, op.numero_wo, op.cliente, os.descricao, os.status, ap.data,
                        ap.horainicio, ap.horafim, imp.nome, et.descricao,
                        ap.giros_rodados, ap.quantidadeproduzida, ap.perdas_producao
                    FROM apontamento AS ap
                    LEFT JOIN ordem_servicos AS os ON ap.servico_id = os.id
                    LEFT JOIN ordem_producao AS op ON os.ordem_id = op.id
                    LEFT JOIN impressores AS imp ON ap.impressor_id = imp.id
                    LEFT JOIN ordem_producao_maquinas AS opm ON os.maquina_id = opm.id
                    LEFT JOIN equipamentos_tipos AS et ON opm.equipamento_id = et.id
                    WHERE {where_clauses}
                    ORDER BY ap.data DESC, ap.horainicio DESC
                    LIMIT %s OFFSET %s
                """

                filters, params = self.get_filters()
                where_str = " AND ".join(filters) if filters else "1=1"
                final_query = query.format(where_clauses=where_str)

                # Calcula o OFFSET e adiciona os novos parâmetros
                offset = (self.current_page - 1) * self.results_per_page
                params.extend([self.results_per_page, offset])

                cur.execute(final_query, tuple(params))
                rows = cur.fetchall()

                # Habilita/desabilita os botões de paginação
                self.prev_page_button.config(state=NORMAL if self.current_page > 1 else DISABLED)
                self.next_page_button.config(state=NORMAL if len(rows) == self.results_per_page else DISABLED)

                for row in rows:
                    processed_row = [item if item is not None else "" for item in row]
                    self.tree_hist.insert("", END, values=tuple(processed_row))
               
                self.calculate_and_display_kpis()

        except (psycopg2.Error, ValueError) as e:
            messagebox.showerror("Erro ao Carregar Histórico", f"Falha ao carregar relatório: {e}", parent=self)
        finally:
            if conn: conn.close()

    def load_pending_orders(self):
        for item in self.tree_pend.get_children():
            self.tree_pend.delete(item)

        conn = self.get_db_connection()
        if not conn: return
        
        try:
            with conn.cursor() as cur:
                query = """
                    SELECT
                        op.sequencia_producao, op.numero_wo, op.cliente,
                        os.descricao, os.status, op.data_previsao_entrega
                    FROM ordem_producao AS op
                    LEFT JOIN ordem_servicos AS os ON op.id = os.ordem_id
                    WHERE (os.status = 'Pendente' OR os.status = 'Em Produção') AND {where_clauses}
                    ORDER BY op.sequencia_producao, os.sequencia
                """
                filters, params = self.get_filters(ignore_dates=True)
                where_str = " AND ".join(filters) if filters else "1=1"
                final_query = query.format(where_clauses=where_str)

                cur.execute(final_query, tuple(params))
                for row in cur.fetchall():
                    processed_row = [item if item is not None else "" for item in row]
                    self.tree_pend.insert("", END, values=tuple(processed_row))

        except (psycopg2.Error, ValueError) as e:
            messagebox.showerror("Erro ao Carregar Pendentes", f"Falha ao carregar ordens pendentes: {e}", parent=self)
        finally:
            if conn: conn.close()

    def get_filters(self, ignore_dates=False):
        filters = []
        params = []
        
        if not ignore_dates:
            if self.date_start_entry.entry.get():
                filters.append("ap.data >= %s")
                params.append(datetime.strptime(self.date_start_entry.entry.get(), '%d/%m/%Y').date())
            if self.date_end_entry.entry.get():
                filters.append("ap.data <= %s")
                params.append(datetime.strptime(self.date_end_entry.entry.get(), '%d/%m/%Y').date())
        
        if self.status_filter.get() != "Todos":
            filters.append("op.status = %s")
            params.append(self.status_filter.get())
        if self.client_filter.get() != "Todos":
            filters.append("op.cliente = %s")
            params.append(self.client_filter.get())
        if self.wo_filter.get() != "Todas":
            filters.append("op.numero_wo = %s")
            params.append(self.wo_filter.get())
            
        return filters, params

    def clear_filters(self):
        self.date_start_entry.entry.delete(0, END)
        self.date_end_entry.entry.delete(0, END)
        self.status_filter.set("Todos")
        self.client_filter.set("Todos")
        self.wo_filter.set("Todos")
        self.apply_all_filters()
        
    def calculate_and_display_kpis(self):
        total_produzido = 0
        total_giros = 0
        total_perdas = 0
        total_producao_delta = timedelta()

        for item_id in self.tree_hist.get_children():
            values = self.tree_hist.item(item_id, 'values')
            idx_inicio = self.cols_hist.index('horainicio')
            idx_fim = self.cols_hist.index('horafim')
            idx_qtd = self.cols_hist.index('qtd_produzida')
            idx_giros = self.cols_hist.index('giros_rodados')
            idx_perdas = self.cols_hist.index('perdas_producao')

            try: total_produzido += int(values[idx_qtd])
            except (ValueError, IndexError): pass
            try: total_giros += int(values[idx_giros])
            except (ValueError, IndexError): pass
            try: total_perdas += int(values[idx_perdas])
            except (ValueError, IndexError): pass
            
            try:
                start_time = datetime.strptime(values[idx_inicio], '%H:%M:%S').time()
                end_time = datetime.strptime(values[idx_fim], '%H:%M:%S').time()
                total_producao_delta += (datetime.combine(date.min, end_time) - datetime.combine(date.min, start_time))
            except (ValueError, IndexError): pass
        
        total_parada_delta = self.calculate_total_downtime()

        total_seconds_prod = int(total_producao_delta.total_seconds())
        h_prod, rem_prod = divmod(total_seconds_prod, 3600)
        m_prod, s_prod = divmod(rem_prod, 60)

        total_seconds_parada = int(total_parada_delta.total_seconds())
        h_parada, rem_parada = divmod(total_seconds_parada, 3600)
        m_parada, s_parada = divmod(rem_parada, 60)

        self.kpi_labels["qtd_produzida"].config(text=f"{total_produzido:,}".replace(",", "."))
        self.kpi_labels["giros_rodados"].config(text=f"{total_giros:,}".replace(",", "."))
        self.kpi_labels["total_perdas"].config(text=f"{total_perdas:,}".replace(",", "."))
        self.kpi_labels["tempo_producao"].config(text=f"{h_prod:02d}:{m_prod:02d}:{s_prod:02d}")
        self.kpi_labels["tempo_parada"].config(text=f"{h_parada:02d}:{m_parada:02d}:{s_parada:02d}")
    
    def calculate_total_downtime(self):
        apontamento_ids = [self.tree_hist.item(item_id, 'values')[0] for item_id in self.tree_hist.get_children()]
        if not apontamento_ids:
            return timedelta()
        
        total_downtime = timedelta()
        conn = self.get_db_connection()
        if not conn: return total_downtime
        try:
            with conn.cursor() as cur:
                placeholders = ','.join(['%s'] * len(apontamento_ids))
                query = f"""
                    SELECT hora_inicio_parada, hora_fim_parada 
                    FROM paradas 
                    WHERE apontamento_id IN ({placeholders})
                """
                cur.execute(query, tuple(apontamento_ids))
                for start_time, end_time in cur.fetchall():
                    if start_time and end_time:
                        total_downtime += (datetime.combine(date.min, end_time) - datetime.combine(date.min, start_time))
        except psycopg2.Error as e:
            print(f"Erro ao calcular tempo de parada: {e}")
        finally:
            if conn: conn.close()
        return total_downtime

    def view_stops(self):
        selected_item = self.tree_hist.focus()
        if not selected_item:
            messagebox.showwarning("Seleção Necessária", "Por favor, selecione um apontamento na lista para ver os detalhes.", parent=self)
            return

        item_values = self.tree_hist.item(selected_item, "values")
        apontamento_id = item_values[0]
        wo_num = item_values[1]
        
        StopsDetailsWindow(self, self.db_config, apontamento_id, wo_num)

    def export_to_xlsx(self):
        if not self.tree_hist.get_children():
            messagebox.showinfo("Exportar", "Não há dados no histórico de produção para exportar.", parent=self)
            return

        filepath = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Planilhas Excel", "*.xlsx"), ("Todos os Arquivos", "*.*")],
            title="Salvar Relatório de Histórico como XLSX"
        )
        if not filepath: return

        try:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Historico_Producao"
            sheet.append(self.headers_hist)
            
            for item_id in self.tree_hist.get_children():
                row_values = self.tree_hist.item(item_id, 'values')
                sheet.append(list(row_values))
                
            workbook.save(filepath)
            messagebox.showinfo("Sucesso", f"Relatório exportado com sucesso para:\n{filepath}", parent=self)
        except Exception as e:
            messagebox.showerror("Erro ao Exportar", f"Ocorreu um erro ao exportar o relatório: {e}", parent=self)

    def open_pcp_window(self):
        # Esta função assume que a janela principal (master) tem um método para abrir a janela do PCP
        # Se a estrutura for diferente, você precisará ajustar.
        # Baseado em main_menu_window.py, a janela mestre deve ter este método.
        if hasattr(self.master, 'open_pcp_window'):
            self.master.open_pcp_window()
        else:
            messagebox.showinfo("Informação", "Função para abrir o PCP não encontrada no contexto atual.")

    def next_page(self):
        self.current_page += 1
        self.load_production_history()

    def prev_page(self):
        if self.current_page > 1:
            self.current_page -= 1
            self.load_production_history()

class StopsDetailsWindow(Toplevel):
    def __init__(self, master, db_config, apontamento_id, wo_number):
        super().__init__(master)
        self.db_config = db_config
        self.apontamento_id = apontamento_id
        
        self.title(f"Detalhes de Parada - WO: {wo_number} (Apont. ID: {apontamento_id})")
        self.geometry("600x400")
        self.grab_set()

        self.create_widgets()
        self.load_stops_data()

    def get_db_connection(self):
        return self.master.get_db_connection()

    def create_widgets(self):
        main_frame = tb.Frame(self, padding=10)
        main_frame.pack(fill=BOTH, expand=True)

        cols = ("motivo", "inicio", "fim", "duracao")
        headers = ("Motivo da Parada", "Hora Início", "Hora Fim", "Duração")
        self.tree = tb.Treeview(main_frame, columns=cols, show="headings", bootstyle=INFO)
        for col, header in zip(cols, headers):
            self.tree.heading(col, text=header)
            self.tree.column(col, anchor=CENTER)
        self.tree.column("motivo", anchor=W)
        self.tree.pack(fill=BOTH, expand=True)

    def load_stops_data(self):
        conn = self.get_db_connection()
        if not conn: return
        try:
            with conn.cursor() as cur:
                query = """
                    SELECT 
                        mt.descricao, 
                        p.hora_inicio_parada, 
                        p.hora_fim_parada
                    FROM paradas p
                    LEFT JOIN motivos_parada_tipos mt ON p.motivo_id = mt.id
                    WHERE p.apontamento_id = %s
                """
                cur.execute(query, (self.apontamento_id,))
                rows = cur.fetchall()

                if not rows:
                    messagebox.showinfo("Sem Paradas", "Nenhuma parada registrada para este apontamento.", parent=self)
                    self.destroy()
                    return

                for motivo, inicio, fim in rows:
                    duracao_str = ""
                    if inicio and fim:
                        delta = datetime.combine(date.min, fim) - datetime.combine(date.min, inicio)
                        duracao_str = str(delta)
                    self.tree.insert("", END, values=(motivo, inicio, fim, duracao_str))
        except psycopg2.Error as e:
            messagebox.showerror("Erro", f"Falha ao carregar detalhes das paradas: {e}", parent=self)
        finally:
            if conn: conn.close()